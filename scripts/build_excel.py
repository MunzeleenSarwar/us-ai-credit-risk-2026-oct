"""Excel model for US AI credit risk."""
import pandas as pd, pathlib
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows
base = pathlib.Path(__file__).resolve().parents[1]
wb = Workbook(); wb.remove(wb.active)
head_fill = PatternFill("solid", fgColor="1F4E79"); head_font = Font(color="FFFFFF", bold=True)
for sheet, path in [("Latest", base/"data"/"latest.csv"), ("Models", base/"data"/"models.csv"), ("Trend", base/"data"/"metrics.csv"), ("Financials", base/"data"/"financials.csv")]:
    df = pd.read_csv(path); ws = wb.create_sheet(sheet)
    for r in dataframe_to_rows(df, index=False, header=True): ws.append(r)
    for c in ws[1]: c.fill = head_fill; c.font = head_font
    ws.freeze_panes = "A2"; ws.sheet_properties.pageSetUpPr.fitToPage = True
ws = wb.create_sheet("Method")
ws.append(["Item", "Detail"])
for row in [
 ["Universe", "10 US-listed AI companies: NVDA, MSFT, GOOGL, META, AVGO, TSM, AMD, PLTR, ORCL, ASML. TSMC is Taiwan based and ASML is Netherlands based, both listed in the US."],
 ["Window", "5 fiscal years per company from yfinance annual statements. Fiscal year ends differ."],
 ["Currency", "Money converted to USD. TSMC statements in TWD converted at 32.562 TWD per USD, ASML in EUR at 1.1178 USD per EUR, FX closes 26 Oct 2026. Ratios do not depend on currency."],
 ["Debt divided by EBITDA", "Total debt divided by EBITDA. Lower means debt is smaller compared with earnings before interest, tax, depreciation and amortisation."],
 ["Interest coverage", "EBIT divided by interest expense. Higher means earnings cover interest more times."],
 ["Altman Z", "1.2 x working capital / assets + 1.4 x retained earnings / assets + 3.3 x EBIT / assets + 0.6 x market value of equity / liabilities + 1.0 x sales / assets. Market value uses current market capitalisation. Scores above 2.99 are usually read as safe, 1.81 to 2.99 as grey, below 1.81 as distress. This score was built for manufacturers and is weak for cash rich technology companies, so it is not used alone for the ranking."],
 ["Ranking", "Ranking is a judgement from debt divided by EBITDA, interest coverage, cash compared with debt, free cash flow compared with debt, and profit history. It is not a credit rating and not a default prediction."],
 ["Sources", "yfinance annual balance sheet, income statement and cash flow statements, plus market capitalisation from yfinance."],
]: ws.append(row)
for c in ws[1]: c.fill = head_fill; c.font = head_font
ws.column_dimensions["A"].width = 28; ws.column_dimensions["B"].width = 130
wb.save(base/"US_AI_Credit_Risk_Model.xlsx")
print("excel done")
